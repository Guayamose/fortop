"""Build the Fortop pitch incrementally, using the original vector brand assets.

Requires python-pptx, cairosvg and fonttools. Run from any directory.
The Windows Bahnschrift font matches the supplied pitch references; set
FORTOP_PITCH_FONT if its installation is elsewhere. Fonts are not redistributed.
"""

from copy import deepcopy
from io import BytesIO
import base64
import os
from pathlib import Path
from xml.etree import ElementTree as ET
from xml.sax.saxutils import escape

import cairosvg
from fontTools.pens.svgPathPen import SVGPathPen
from fontTools.ttLib import TTFont
from fontTools.varLib.instancer import instantiateVariableFont
from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_AUTO_SIZE, MSO_ANCHOR
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[2]
PITCH = ROOT / "fortop_pitch"
BRAND = ROOT / "fortop_brand_kit"
SLIDES = PITCH / "slides"
ASSETS = PITCH / "assets" / "01-portada"
W, H = 1920, 1080
PALETTE = {"shadow": "#2E282A", "pine": "#2B4F42", "ivory": "#FFFFF3", "honeydew": "#E0EFDA"}
SVG_NS = "http://www.w3.org/2000/svg"
ASVG_NS = "http://schemas.microsoft.com/office/drawing/2016/SVG/main"
FONT_PATH = Path(os.environ.get("FORTOP_PITCH_FONT", "/mnt/c/Windows/Fonts/bahnschrift.ttf"))
FONT_FAMILY = "Bahnschrift"
CODE_FONT_PATH = Path(os.environ.get("FORTOP_PITCH_CODE_FONT", "/mnt/c/Windows/Fonts/consola.ttf"))
DEMO_URL = "https://fortop.34-76-10-0.sslip.io"
TEAM_MEMBERS = (
    "Rafael Pradillo",
    "Álvaro Fernández-Aller",
    "Carlos Blázquez",
    "Alejandro Ferreras",
    "Jaime Poyatos",
)
TEAM_CREDITS = "  ·  ".join(TEAM_MEMBERS)
TEXT = [
    {"text": "Visibilidad de marca", "x": 132, "baseline": 640, "size": 48, "color": "shadow", "weight": 400},
    {"text": "y remediación SEO/GEO.", "x": 132, "baseline": 706, "size": 48, "color": "pine", "weight": 400},
    {"text": TEAM_CREDITS, "x": 132, "baseline": 966, "size": 26, "color": "pine", "weight": 400, "width": 1656},
]
NOTES = """PORTADA · 10 segundos

Frase aprobada: Visibilidad de marca y remediación SEO/GEO.

Fortop analiza la visibilidad de una marca y utiliza agentes para detectar y corregir problemas SEO/GEO en su web.

Transición: empezar por la situación del cliente, antes de explicar el producto.
Formato acordado: 3 minutos de pitch en español y 1 minuto adicional de demo.
"""

PROBLEM_TEXT = [
    {"text": "EL PROBLEMA", "x": 132, "baseline": 112, "size": 25, "color": "pine", "weight": 400, "width": 480},
    {"text": "Tus clientes piden recomendaciones.", "x": 132, "baseline": 232, "size": 68, "color": "shadow", "weight": 400, "width": 1656},
    {"text": "¿Aparece tu marca?", "x": 132, "baseline": 318, "size": 68, "color": "pine", "weight": 400, "width": 1656},
    {"text": "¿Qué software de gestión", "x": 174, "baseline": 568, "size": 41, "color": "shadow", "weight": 400, "width": 548},
    {"text": "recomiendas para una pyme?", "x": 174, "baseline": 626, "size": 41, "color": "shadow", "weight": 400, "width": 548},
    {"text": "01", "x": 1234, "baseline": 513, "size": 24, "color": "pine", "weight": 400, "width": 55},
    {"text": "Competidor A", "x": 1302, "baseline": 513, "size": 35, "color": "shadow", "weight": 400, "width": 486},
    {"text": "02", "x": 1234, "baseline": 587, "size": 24, "color": "pine", "weight": 400, "width": 55},
    {"text": "Competidor B", "x": 1302, "baseline": 587, "size": 35, "color": "shadow", "weight": 400, "width": 486},
    {"text": "03", "x": 1234, "baseline": 661, "size": 24, "color": "pine", "weight": 400, "width": 55},
    {"text": "Competidor C", "x": 1302, "baseline": 661, "size": 35, "color": "shadow", "weight": 400, "width": 486},
    {"text": "¿Y tu marca?", "x": 1302, "baseline": 763, "size": 38, "color": "pine", "weight": 400, "width": 486},
    {"text": "¿Dónde apareces?", "x": 189, "baseline": 930, "size": 34, "color": "shadow", "weight": 400, "width": 465},
    {"text": "¿A quién recomiendan?", "x": 767, "baseline": 930, "size": 34, "color": "shadow", "weight": 400, "width": 465},
    {"text": "¿Qué debes corregir?", "x": 1343, "baseline": 930, "size": 34, "color": "shadow", "weight": 400, "width": 445},
    {"text": "02 / 08", "x": 1708, "baseline": 1022, "size": 20, "color": "pine", "weight": 400, "width": 80},
]

PROBLEM_NOTES = """PROBLEMA · 25 segundos

Una empresa puede tener una web cuidada y seguir quedando fuera cuando un cliente pide una recomendación. Necesita entender en qué consultas aparece, a qué competidores recomiendan y qué problemas de contenido o código debe corregir. Eso exige investigar respuestas, contrastar fuentes y coordinar cambios en la web.

El ejemplo de software para pymes y los competidores A, B y C son ilustrativos. No representan una respuesta medida ni clientes reales. No se presentan cifras de adopción ni pérdidas de ingresos.

Los símbolos de Grok, Gemini, Claude y OpenAI identifican asistentes en los que un cliente podría pedir recomendaciones. No representan integraciones disponibles en Fortop: el estudio actual mide Gemini API con búsqueda habilitada. Los SVG proceden de LobeHub Icons; originales, licencia y fuentes se conservan en assets/02-problema/logos.
https://github.com/lobehub/lobe-icons

Transición al insight: la visibilidad depende del contexto de la marca, del contenido y de la base técnica de su web; entender el problema y ejecutar mejoras son partes del mismo trabajo.
"""

INSIGHT_NOTES = """INSIGHT · 20 segundos

Tener una web no basta. Hay dos capas que influyen en su visibilidad: lo que cuenta y cómo está construida. El contenido debe responder a lo que busca el cliente. La base técnica permite que los buscadores rastreen e interpreten esa información. Una buena oferta puede perder oportunidades si una de esas capas falla.

QUÉ EXPLICA LA SLIDE
El insight presenta dos dimensiones complementarias del problema: relevancia del contenido y acceso técnico a la información publicada. No presenta el producto, sus agentes ni sus funcionalidades. La solución se introduce por primera vez en la diapositiva siguiente.

La relevancia depende de la necesidad del público y del contenido que la responde. El acceso técnico incluye permitir el rastreo de las páginas públicas y facilitar su interpretación mediante texto, enlaces y metadatos coherentes.

Estas dimensiones no son una lista exhaustiva de factores ni garantizan aparecer en una recomendación. También influyen las fuentes externas, la consulta y los sistemas de cada proveedor. Una marca puede ser mencionada a partir de fuentes de terceros aunque su web no sea rastreable. «Ser encontrado» se refiere aquí a facilitar el descubrimiento del contenido propio.

BASE TÉCNICA Y FUENTES
Google indica que las buenas prácticas SEO siguen siendo aplicables a sus funciones de búsqueda generativa: acceso al rastreo, contenido textual, enlaces internos y datos estructurados coherentes. No exige archivos ni marcado especiales, y cumplir los requisitos no garantiza aparición.
https://developers.google.com/search/docs/appearance/ai-features

Su guía de contenido útil recomienda atender al público previsto, aportar información fiable y valorar la experiencia de la página. La organización de estas ideas en dos capas es una síntesis para el pitch, no una clasificación oficial ni una promesa de posicionamiento.
https://developers.google.com/search/docs/fundamentals/creating-helpful-content

Fuentes consultadas el 22 de octubre de 2026. Transición: Fortop reúne investigación de marca y remediación técnica en una plataforma.
"""

SOLUTION_NOTES = """SOLUCIÓN · 25 segundos

Fortop trabaja esas dos capas desde una plataforma. Subes un brief o documentos de tu marca y recibes un estudio de presencia, competencia y fuentes, con evidencias y acciones priorizadas. O conectas tu repositorio de GitHub: los agentes analizan el código, proponen correcciones y publican issues y pull requests para que los revises. El informe y el historial quedan guardados en la plataforma.

ALCANCE DE LA ILUSTRACIÓN
Las dos vistas son ilustraciones simplificadas del producto, no capturas ni resultados de un estudio real. Las barras no representan porcentajes medidos. El cambio de título HTML es un ejemplo de corrección, no un cambio pendiente ni publicado en un repositorio real.

El estudio actual mide Gemini API con búsqueda habilitada; no se presentan mediciones de ChatGPT o Claude. Los informes incluyen respuestas, citas, competencia, mercados y un plan de acción. La investigación de marca y el análisis de repositorios son recorridos independientes; una recomendación técnica puede enviarse a un repositorio conectado.

Las correcciones pasan por comprobaciones y se proponen en GitHub; no se fusionan ni despliegan automáticamente. Un análisis también puede concluir sin cambios publicables. El contenido editorial y las acciones externas no se resuelven automáticamente al abrir un PR.

BASE EN EL PRODUCTO
fortop_platform/app/views/brands/_report.html.erb
fortop_platform/app/views/brands/_document_upload.html.erb
fortop_platform/app/views/repositories/_analysis.html.erb

Transición: explicar cómo los agentes investigan, corrigen y verifican antes de entregar los resultados.
"""

AGENTS_NOTES = """SISTEMA AGÉNTICO · 35 segundos

Detrás hay Google Gemini y herramientas propias. En marca, investigamos la web, repetimos preguntas por mercado y contrastamos identidades y citas antes de priorizar acciones. En código, un agente analiza y entrega hallazgos a otro que edita los archivos. El resultado de cada herramienta vuelve al agente: si una comprobación falla, puede reparar y repetir. Después, una verificación independiente controla la publicación. Los cambios llegan en un PR para revisión; las acciones y los resultados quedan guardados.

CÓMO LEER EL DIAGRAMA
Las dos filas resumen recorridos independientes. Los nodos son etapas del sistema, no ocho agentes autónomos ni ocho modelos. Se usa Google Gemini en ambos verticales. Las comprobaciones deterministas y la publicación las ejecuta la aplicación.

MARCA
El flujo interpreta el brief, investiga con búsqueda web, diseña preguntas por mercado y recoge respuestas independientes repetidas. Las preguntas de descubrimiento no incluyen el brief ni la marca esperada. Después extrae identidades por contexto y contrasta citas; las métricas se calculan en código. El diagnóstico vincula sus acciones a observaciones guardadas. «Investiga» agrupa interpretación e investigación; «Mide» incluye diseño del panel y mediciones. La fila no presenta un sistema de agentes que elija libremente todas las etapas: es un flujo orquestado con llamadas especializadas.

REPOSITORIO
El agente de análisis elige herramientas para listar, leer y buscar archivos y registrar hallazgos respaldados por citas del código leído. Un segundo agente recibe los hallazgos automatizables y utiliza herramientas de edición y validación. El runner devuelve el resultado de cada herramienta al modelo para que elija su siguiente acción.

La flecha «Si falla, corrige» representa validate_changes dentro del ciclo de remediación. Si no puede preparar un cambio seguro, puede descartarlo y marcar el hallazgo para intervención manual. Al terminar, la aplicación hace otra verificación independiente antes de publicar. No es un reintento automático ilimitado tras cualquier fallo final.

La validación comprueba alcance, sintaxis y algunas condiciones semánticas, según el formato. No sustituye la suite de tests ni el CI del repositorio. Las propuestas se publican en una rama separada, con issues y PR cuando corresponde; no se fusionan ni despliegan automáticamente.

BASE EN EL CÓDIGO
fortop_platform/app/services/visibility/pipeline.rb
fortop_platform/app/services/visibility/evidence.rb
fortop_platform/app/services/visibility/metrics.rb
fortop_platform/app/services/agents/runner.rb
fortop_platform/app/services/agents/analysis_tools.rb
fortop_platform/app/services/agents/remediation_tools.rb
fortop_platform/app/services/scans/pipeline.rb
fortop_platform/app/services/scans/validator.rb

Transición: enseñar la evidencia del producto construido y los resultados revisables.
"""

PRODUCT_NOTES = f"""DEMO · 25 segundos de presentación; demostración de 1 minuto aparte

Vamos a verlo en la aplicación. Primero, un estudio de marca con sus respuestas y fuentes. Después, un hallazgo del repositorio, su corrección y la propuesta en GitHub. El informe y los cambios quedan disponibles para revisarlos.

ABRIR LA APLICACIÓN
{DEMO_URL}
El botón «Abrir demo» enlaza a la web en PowerPoint y PDF. La web requiere iniciar sesión con una cuenta autorizada. Para presentar, deja preparada la sesión de GitHub en el navegador.

RECORRIDO PROPUESTO · 60 SEGUNDOS
0–25 s: abrir un estudio de marca ya completado, mostrar el informe y una respuesta con sus fuentes.
25–50 s: abrir un análisis de repositorio ya completado, enseñar un hallazgo y su diff, y mostrar el enlace al PR.
50–60 s: volver a la presentación.

Esta diapositiva sustituye las capturas de producto. Las imágenes y evidencias anteriores se conservan en assets/06-producto como material de trabajo y ya no aparecen en la slide. El minuto de demostración es adicional a los tres minutos de exposición acordados.
"""

BUSINESS_NOTES = """CLIENTE Y NEGOCIO · 25 segundos

El cliente inicial es el fundador de un SaaS pequeño, con web en GitHub y sin especialista SEO. Fortop analiza la presencia de su marca en Gemini y prepara correcciones para su web que su desarrollador puede revisar. Proponemos una cuota mensual según el volumen de análisis. Primero lo probaríamos con fundadores para validar la utilidad y su disposición a pagar.

MODELO PROPUESTO, PENDIENTE DE VALIDACIÓN COMERCIAL
La diapositiva describe el segmento inicial acordado para el pitch: pequeñas empresas de software por suscripción, sin especialista SEO y con la web en GitHub. No afirma que existan clientes de pago, ingresos, contratos ni pilotos en marcha. No fija precios, ahorro, margen, conversión ni tamaño de mercado. La cuota mensual es una hipótesis de ingresos, no una tarifa comercial ya disponible.

CLIENTE INICIAL
Comprador: fundador de un SaaS pequeño, que necesita captar clientes y no tiene a una persona dedicada al SEO. Requisitos para probar ambos verticales: web propia en GitHub, acceso autorizado al repositorio y una persona capaz de revisar y aprobar los cambios. SaaS significa software por suscripción.

Ejemplo hipotético para explicar el perfil si se pregunta: una empresa vende software de reservas para academias. Tiene producto y web, pero necesita entender si la recomiendan, frente a quién compite y qué puede mejorar en su sitio. No es un cliente real ni un caso de resultados demostrados.

Los tres iconos representan al fundador, el diagnóstico con propuestas de código y el cobro mensual. El estudio de marca también puede utilizarse sin repositorio; la selección inicial de webs en GitHub permite validar la propuesta completa.

INGRESOS
La empresa contrataría una suscripción mensual para analizar su marca y su web. El precio variaría según el volumen de análisis: estudios de marca y análisis con remediación de código. Los precios y las condiciones comerciales están pendientes de diseño y prueba; no describen límites actuales de los agentes ni un sistema de cobro ya implementado. No se afirma que exista programación automática de análisis recurrentes.

SERVICIO
El diagnóstico mide presencia, competidores y fuentes en Gemini API con búsqueda habilitada. No mide las interfaces nativas de Gemini, ChatGPT o Claude. Las correcciones web se entregan como propuestas para revisión en GitHub; el fundador y su desarrollador conservan la decisión de fusionarlas y desplegarlas. No se promete que Fortop resuelva automáticamente toda acción editorial ni garantice posicionamiento.

ENTRADA AL MERCADO
Pruebas acompañadas con fundadores del perfil descrito. Evaluar utilidad del diagnóstico, aceptación de correcciones, tiempo de trabajo y coste de ejecución. Contrastar esas métricas con disposición a pagar antes de escalar. El ahorro es una variable a medir, no un resultado ya demostrado.

BASE DEL MODELO
Segmento y comprador definidos con el usuario para este pitch. Suscripción mensual según volumen como hipótesis comercial. El producto actual ofrece los dos verticales mostrados en las slides anteriores; la validación de mercado está pendiente.

Transición: cerrar con la visión de Fortop.
"""

CLOSING_NOTES = """CIERRE · 15 segundos

Fortop une el diagnóstico de tu marca con mejoras concretas para tu web. Queremos que una pequeña empresa pueda entender su visibilidad y actuar sobre ella, aunque no tenga equipo SEO. Gracias.

IDEA DE CIERRE
Entender la visibilidad de marca y poder actuar sobre la web desde un mismo producto. La frase resume los dos verticales construidos y el propósito para el cliente inicial. No introduce funciones futuras, cifras ni promesas de posicionamiento. Las correcciones siguen siendo propuestas que una persona revisa antes de fusionar y desplegar.

REPARTO PREVISTO DEL PITCH
Portada: 10 s. Problema: 25 s. Insight: 20 s. Solución: 25 s.
Agentes: 35 s. Producto: 25 s. Cliente y negocio: 25 s. Cierre: 15 s.
Total: 180 segundos. Demo aparte: 60 segundos.
"""


def svg_document(body, width=W, height=H, title="Fortop · portada"):
    return (f'<svg xmlns="{SVG_NS}" width="{width}" height="{height}" viewBox="0 0 {width} {height}" role="img">'
            f'<title>{title}</title>{body}</svg>')


def group_from_asset(relative):
    root = ET.parse(BRAND / relative).getroot()
    group = deepcopy(root.find(f"{{{SVG_NS}}}g"))
    return ET.tostring(group, encoding="unicode").replace("ns0:", "").replace(":ns0", "")


def logo_asset():
    # The original asset has a 24-unit safety margin. Only remove that margin
    # from its viewport; preserve every original path, color and proportion.
    logo = group_from_asset("03_IMAGOTIPO/FORTOP_imagotipo_color-transparente.svg")
    return (f'<svg xmlns="{SVG_NS}" width="864.8" height="192" viewBox="24 24 864.8 192">'
            f'<title>FORTOP · imagotipo original</title>{logo}</svg>')


def ascenso_art():
    # A shallow architectural relief built from Ascenso's original four planes.
    # Keep the channels open and shade only faces visible along the extrusion.
    # This decorative asset never replaces the original foreground brand mark.
    planes = [
        [(0, 142), (48, 88), (48, 172), (0, 192)],
        [(64, 66), (110, 0), (110, 146), (64, 165)],
        [(122, 16), (164, 68), (164, 165), (122, 147)],
        [(178, 88), (224, 142), (224, 192), (178, 172)],
    ]
    dx, dy = 8, 6

    def points(vertices):
        return " ".join(f"{x:.3f},{y:.3f}" for x, y in vertices)

    def mix(first, second, amount):
        a = [int(first[index:index + 2], 16) for index in (1, 3, 5)]
        b = [int(second[index:index + 2], 16) for index in (1, 3, 5)]
        return "#" + "".join(f"{round(x * (1 - amount) + y * amount):02x}" for x, y in zip(a, b))

    ivory, honeydew, pine = (PALETTE[key] for key in ("ivory", "honeydew", "pine"))
    defs = f'''<defs>
      <linearGradient id="ascenso-front" gradientUnits="userSpaceOnUse" x1="10" y1="0" x2="240" y2="195">
        <stop offset="0" stop-color="{ivory}"/>
        <stop offset=".48" stop-color="{mix(ivory, honeydew, .38)}"/>
        <stop offset="1" stop-color="{mix(ivory, honeydew, .92)}"/>
      </linearGradient>
    </defs>'''

    # Dense, low-opacity contours approximate soft contact shadows as vectors.
    # No stepped translated silhouettes or renderer-dependent blur filters.
    shadows = []
    for vertices in planes:
        shifted = [(x + dx + 3.5, y + dy + 4) for x, y in vertices]
        shadows.append(f'<polygon points="{points(shifted)}" fill="{pine}" opacity=".025"/>')
        for index in range(36, 0, -1):
            radius = index * .32
            alpha = .0006 + .0016 * (1 - index / 36) ** 2
            shadows.append(f'<polygon points="{points(shifted)}" fill="none" stroke="{pine}" '
                           f'stroke-width="{radius * 2:.3f}" stroke-linejoin="round" opacity="{alpha:.5f}"/>')

    solids = []
    for vertices in reversed(planes):
        sides, edges = [], []
        for index, a in enumerate(vertices):
            b = vertices[(index + 1) % len(vertices)]
            edge_x, edge_y = b[0] - a[0], b[1] - a[1]
            length = (edge_x ** 2 + edge_y ** 2) ** .5
            normal_x, normal_y = edge_y / length, -edge_x / length
            if normal_x * dx + normal_y * dy > 0:
                face = [a, b, (b[0] + dx, b[1] + dy), (a[0] + dx, a[1] + dy)]
                shade = .115 + .035 * max(0, normal_y)
                sides.append(f'<polygon points="{points(face)}" fill="{mix(ivory, pine, shade)}"/>')
                sides.append(f'<path d="M{a[0] + dx} {a[1] + dy} L{b[0] + dx} {b[1] + dy}" '
                             f'fill="none" stroke="{pine}" stroke-width=".28" stroke-opacity=".08"/>')
            light = -.65 * normal_x - .76 * normal_y
            color = "#FFFFFF" if light > 0 else pine
            opacity = .82 if light > 0 else .12
            edges.append(f'<path d="M{a[0]} {a[1]} L{b[0]} {b[1]}" fill="none" '
                         f'stroke="{color}" stroke-width=".6" stroke-opacity="{opacity}" stroke-linecap="round"/>')
        solids += sides
        solids.append(f'<polygon points="{points(vertices)}" fill="url(#ascenso-front)"/>')
        solids += edges
    return svg_document(defs + '<g transform="translate(82 122) scale(2.85)">' +
                        "".join(shadows + solids) + '</g>', 860, 760,
                        "Ascenso · planos separados en relieve de marfil")


_fonts = {}


def font_for(weight, font_path=FONT_PATH):
    key = (weight, str(font_path))
    if key not in _fonts:
        font = TTFont(font_path)
        if "fvar" in font:
            axes = {axis.axisTag: axis.defaultValue for axis in font["fvar"].axes}
            if "wght" in axes:
                axes["wght"] = weight
            font = instantiateVariableFont(font, axes, inplace=True)
        _fonts[key] = font
    return _fonts[key]


def text_paths(item):
    if item.get("segments"):
        x, paths = item["x"], []
        for segment in item["segments"]:
            part = {**item, **segment, "x": x}
            part.pop("segments")
            outlines, width = text_paths(part)
            paths.append(outlines)
            x += width
        return "".join(paths), x - item["x"]
    font = font_for(item["weight"], item.get("font_path", FONT_PATH))
    glyphs = font.getGlyphSet()
    cmap = font.getBestCmap()
    scale = item["size"] / font["head"].unitsPerEm
    x = item["x"]
    paths = []
    for char in item["text"]:
        glyph_name = cmap[ord(char)]
        pen = SVGPathPen(glyphs)
        glyphs[glyph_name].draw(pen)
        if pen.getCommands():
            paths.append(f'<path d="{pen.getCommands()}" transform="translate({x:.4f} {item["baseline"]}) '
                         f'scale({scale:.7f} {-scale:.7f})"/>')
        x += font["hmtx"][glyph_name][0] * scale
    return (f'<g fill="{PALETTE[item["color"]]}" aria-label="{escape(item["text"])}">' + "".join(paths) + '</g>', x - item["x"])


def place(svg, x, y, width, height):
    element = ET.fromstring(svg)
    element.attrib.update(x=str(x), y=str(y), width=str(width), height=str(height))
    return ET.tostring(element, encoding="unicode").replace("ns0:", "").replace(":ns0", "")


def add_vector_picture(slide, source, name, x, y, width, height):
    source_bytes = source.encode()
    fallback = cairosvg.svg2png(bytestring=source_bytes, output_width=round(width * 2), output_height=round(height * 2))
    picture = slide.shapes.add_picture(BytesIO(fallback), Inches(x / 144), Inches(y / 144),
                                       width=Inches(width / 144), height=Inches(height / 144))
    picture.name = name
    package = slide.part.package
    svg_part = Part(PackURI(f"/ppt/media/{name}.svg"), "image/svg+xml", package, source_bytes)
    relationship = slide.part.relate_to(svg_part, RT.IMAGE)
    blip = picture._pic.blipFill.blip
    extension_list = etree.SubElement(blip, "{http://schemas.openxmlformats.org/drawingml/2006/main}extLst")
    extension = etree.SubElement(extension_list, "{http://schemas.openxmlformats.org/drawingml/2006/main}ext",
                                  uri="{96DAC541-7B7A-43D3-8B79-37D633B846F1}")
    svg_blip = etree.SubElement(extension, f"{{{ASVG_NS}}}svgBlip", nsmap={"asvg": ASVG_NS})
    svg_blip.set("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed", relationship)
    return picture


def add_native_text(slide, item):
    # A zero-margin text box remains editable in PowerPoint. The companion
    # SVG/PDF use glyph outlines for an exact preview without installed fonts.
    top = item["baseline"] - item["size"] * .95
    box = slide.shapes.add_textbox(Inches(item["x"] / 144), Inches(top / 144),
                                  Inches(item.get("width", 1010) / 144), Inches(item["size"] * 1.5 / 144))
    box.name = item["text"]
    frame = box.text_frame
    frame.clear()
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.word_wrap = False
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.vertical_anchor = MSO_ANCHOR.TOP
    paragraph = frame.paragraphs[0]
    paragraph.space_before = paragraph.space_after = Pt(0)
    paragraph.line_spacing = 1
    for segment in item.get("segments", [{"text": item["text"]}]):
        run = paragraph.add_run()
        run.text = segment["text"]
        run.font.name = item.get("font_family", FONT_FAMILY)
        run.font.size = Pt(item["size"] / 2)
        run.font.bold = segment.get("weight", item["weight"]) >= 600
        run.font.color.rgb = RGBColor.from_string(PALETTE[segment.get("color", item["color"])][1:])
        run._r.get_or_add_rPr().set("lang", "es-ES")


def problem_diagram():
    pine, honeydew = PALETTE["pine"], PALETTE["honeydew"]
    return svg_document(f'''
      <g fill="none" stroke="{pine}" stroke-width="2" stroke-linecap="round" stroke-linejoin="round">
        <!-- A single chat message, centered on the logos and response flow. -->
        <path d="M160 478 H736 Q764 478 764 506 V658 Q764 686 736 686 H234 L210 706 V686 H160 Q132 686 132 658 V506 Q132 478 160 478 Z" fill="{honeydew}" fill-opacity=".38" stroke-opacity=".22" stroke-width="1.5"/>
        <path d="M792 582 H846 M836 576 L846 582 836 588 M1092 582 H1190 M1180 576 L1190 582 1180 588" stroke-opacity=".5" stroke-width="1.8"/>
        <path d="M1230 540 H1788 M1230 614 H1788 M1230 688 H1788" stroke="{honeydew}" stroke-width="1.5"/>
        <!-- The missing brand gets breathing room below the three results. -->
        <circle cx="1251" cy="749" r="19" stroke-opacity=".6" stroke-dasharray="2 5" stroke-width="1.8"/>
        <path d="M1245 743 C1245 735 1258 735 1258 743 C1258 747 1251 748 1251 753 M1251 760 V760.2" stroke-opacity=".85"/>
        <path d="M132 852 H1788" stroke-opacity=".2"/>
        <circle cx="149" cy="918" r="14"/>
        <ellipse cx="149" cy="918" rx="6" ry="14" stroke-width="1.3"/>
        <path d="M135 918 H163" stroke-width="1.3"/>
        <circle cx="724" cy="911" r="7"/>
        <path d="M710 932 V927 C710 916 738 916 738 927 V932 M743 906 C753 906 754 918 747 921 M746 926 C754 927 758 930 758 934"/>
        <path d="M1285 928 L1290 914 1310 894 1320 904 1300 924 1285 928 Z M1290 914 L1300 924 M1305 899 L1315 909"/>
      </g>
    ''', title="Problema · consulta y recomendación ilustrativas")


def problem_assistant_logos():
    # Original vector silhouettes, optical sizing and generous space between
    # marks. Keep each logo independently selectable in the PowerPoint deck.
    placements = [
        ("grok", 906, 522, 70),
        ("gemini", 1026, 522, 74),
        ("claude", 906, 642, 68),
        ("openai", 1026, 642, 70),
    ]
    logos = []
    for name, center_x, center_y, size in placements:
        source = (PITCH / "assets" / "02-problema" / "logos" / f"{name}.svg").read_text()
        root = ET.fromstring(source.replace("currentColor", PALETTE["pine"]))
        root.attrib.pop("style", None)
        root.set("width", "24")
        root.set("height", "24")
        root.set("role", "img")
        svg = ET.tostring(root, encoding="unicode").replace("ns0:", "").replace(":ns0", "")
        logos.append((name, svg, center_x - size / 2, center_y - size / 2, size))
    return logos


def build_problem(deck):
    assets = PITCH / "assets" / "02-problema"
    assets.mkdir(parents=True, exist_ok=True)
    diagram = problem_diagram()
    logos = problem_assistant_logos()
    (assets / "recomendacion-ilustrativa.svg").write_text(diagram)
    symbol_group = group_from_asset("01_ISOTIPO/FORTOP_isotipo_color-transparente.svg")
    symbol = f'<svg xmlns="{SVG_NS}" viewBox="24 24 224 192" width="224" height="192">{symbol_group}</svg>'
    body = f'<rect width="{W}" height="{H}" fill="{PALETTE["ivory"]}"/>'
    body += place(diagram, 0, 0, W, H)
    for name, logo, x, y, size in logos:
        body += place(logo, x, y, size, size)
    body += place(symbol, 1738, 75, 50, 42.86)
    for item in PROBLEM_TEXT:
        outlines, width = text_paths(item)
        if width > item["width"]:
            raise ValueError(f'Text exceeds its column: {item["text"]} ({width:.1f} > {item["width"]})')
        body += outlines
    svg = svg_document(body, title="Fortop · problema")
    (SLIDES / "02-problema.svg").write_text(svg)
    cairosvg.svg2png(bytestring=svg.encode(), write_to=str(SLIDES / "02-problema.png"), output_width=W, output_height=H)
    cairosvg.svg2pdf(bytestring=svg.encode(), write_to=str(SLIDES / "02-problema.pdf"), output_width=1280, output_height=720)

    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string(PALETTE["ivory"][1:])
    add_vector_picture(slide, diagram, "problema-diagrama", 0, 0, W, H)
    for name, logo, x, y, size in logos:
        add_vector_picture(slide, logo, f"problema-logo-{name}", x, y, size, size)
    add_vector_picture(slide, symbol, "problema-isotipo", 1738, 75, 50, 42.86)
    for item in PROBLEM_TEXT:
        add_native_text(slide, item)
    slide.notes_slide.notes_text_frame.text = PROBLEM_NOTES


def insight_diagram():
    ivory, honeydew = PALETTE["ivory"], PALETTE["honeydew"]
    return svg_document(f'''
      <g fill="none" stroke="{honeydew}" stroke-width="2.5" stroke-linecap="round" stroke-linejoin="round" stroke-opacity=".55">
        <path d="M604 448 H738 M725 438 L738 448 725 458 M1182 448 H1316 M1303 438 L1316 448 1303 458"/>
        <path d="M604 772 H738 M725 762 L738 772 725 782 M1182 772 H1316 M1303 762 L1316 772 1303 782"/>
        <path d="M132 664 H1788" stroke-opacity=".17" stroke-width="1.5"/>
      </g>
      <g fill="none" stroke="{ivory}" stroke-width="3.5" stroke-linecap="round" stroke-linejoin="round">
        <!-- Content, the customer's question, and relevance. -->
        <path d="M346 397 H394 L418 421 V495 H346 Z" fill="{honeydew}" fill-opacity=".06"/>
        <path d="M394 397 V421 H418 M363 442 H401 M363 460 H401 M363 478 H388" stroke-opacity=".7"/>
        <path d="M918 406 H1002 Q1013 406 1013 417 V466 Q1013 477 1002 477 H962 L939 495 V477 H918 Q907 477 907 466 V417 Q907 406 918 406 Z" fill="{honeydew}" fill-opacity=".06"/>
        <path d="M930 431 H990 M930 450 H971" stroke-opacity=".7"/>
        <circle cx="1538" cy="448" r="43"/>
        <circle cx="1538" cy="448" r="26" stroke-opacity=".65"/>
        <circle cx="1538" cy="448" r="8" fill="{honeydew}" stroke="none"/>
        <!-- The technical layer is a property of the web, not a product action. -->
        <path d="M352 738 L328 761 352 784 M412 738 L436 761 412 784 M393 724 L371 799"/>
        <!-- Engines can read a web page. -->
        <rect x="908" y="725" width="104" height="83" rx="7" fill="{honeydew}" fill-opacity=".06"/>
        <path d="M908 746 H1012 M924 763 H980 M924 779 H966 M924 795 H952" stroke-opacity=".7"/>
        <!-- Purpose: engines can discover the website. -->
        <circle cx="1529" cy="760" r="36"/>
        <path d="M1555 786 L1588 819" stroke-width="5"/>
        <path d="M1510 753 H1548 M1510 768 H1536" stroke-opacity=".7"/>
      </g>
    ''', title="Dos capas de visibilidad: contenido relevante y base técnica")


def build_insight(deck):
    assets = PITCH / "assets" / "03-insight"
    assets.mkdir(parents=True, exist_ok=True)
    diagram = insight_diagram()
    (assets / "diagnostico-conectado.svg").write_text(diagram)
    symbol_group = group_from_asset("01_ISOTIPO/FORTOP_isotipo_mono-ivory.svg")
    symbol = f'<svg xmlns="{SVG_NS}" viewBox="24 24 224 192" width="224" height="192">{symbol_group}</svg>'

    def label(value, x, baseline, size, color="ivory", width=1656, center=None):
        item = {"text": value, "x": x, "baseline": baseline, "size": size, "color": color, "weight": 400, "width": width}
        if center is not None:
            _, measured = text_paths(item)
            item["x"] = center - measured / 2
            item["width"] = min(width, W - 132 - item["x"])
        return item

    texts = [
        label("EL INSIGHT", 132, 112, 25),
        label("Dos capas de visibilidad.", 132, 232, 68),
        label("LA CAPA", 0, 333, 25, "honeydew", width=470, center=382),
        label("QUÉ PERMITE", 0, 333, 25, "honeydew", width=470, center=960),
        label("POR QUÉ IMPORTA", 0, 333, 25, "honeydew", width=470, center=1538),

        label("Contenido", 0, 587, 40, width=470, center=382),
        label("Responder a", 0, 561, 40, width=470, center=960),
        label("lo que buscan", 0, 614, 40, width=470, center=960),
        label("Ser relevante", 0, 587, 40, width=470, center=1538),

        label("Base técnica", 0, 909, 40, width=470, center=382),
        label("Rastrear y entender", 0, 883, 40, width=470, center=960),
        label("la web", 0, 936, 40, width=470, center=960),
        label("Ser encontrado", 0, 909, 40, width=470, center=1538),
        label("03 / 08", 1708, 1022, 20, "honeydew", width=80),
    ]
    body = f'<rect width="{W}" height="{H}" fill="{PALETTE["pine"]}"/>'
    body += place(diagram, 0, 0, W, H)
    body += place(symbol, 1738, 75, 50, 42.86)
    for item in texts:
        outlines, width = text_paths(item)
        if width > item["width"] or item["x"] < 100 or item["x"] + width > W - 100:
            raise ValueError(f'Text exceeds its column: {item["text"]}')
        body += outlines
    svg = svg_document(body, title="Fortop · insight")
    (SLIDES / "03-insight.svg").write_text(svg)
    cairosvg.svg2png(bytestring=svg.encode(), write_to=str(SLIDES / "03-insight.png"), output_width=W, output_height=H)
    cairosvg.svg2pdf(bytestring=svg.encode(), write_to=str(SLIDES / "03-insight.pdf"), output_width=1280, output_height=720)

    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string(PALETTE["pine"][1:])
    add_vector_picture(slide, diagram, "insight-diagrama", 0, 0, W, H)
    add_vector_picture(slide, symbol, "insight-isotipo", 1738, 75, 50, 42.86)
    for item in texts:
        add_native_text(slide, item)
    slide.notes_slide.notes_text_frame.text = INSIGHT_NOTES


def solution_icon(name, x, y, size=32):
    # One square grid and one stroke weight; icons keep their native proportions.
    paths = {
        "responses": '<path d="M21 14a3 3 0 0 1-3 3H8l-5 4V6a3 3 0 0 1 3-3h12a3 3 0 0 1 3 3Z"/><path d="M7 8h10M7 12h6"/>',
        "sources": '<circle cx="12" cy="12" r="9"/><path d="M3 12h18M12 3c3 3 4 6 4 9s-1 6-4 9c-3-3-4-6-4-9s1-6 4-9Z"/>',
        "evidence": '<path d="M14 2H6a2 2 0 0 0-2 2v16a2 2 0 0 0 2 2h12a2 2 0 0 0 2-2V8Z"/><path d="M14 2v6h6M8 15l3 3 5-6"/>',
        "brief": '<path d="M14 2H6a2 2 0 0 0-2 2v16a2 2 0 0 0 2 2h12a2 2 0 0 0 2-2V8Z"/><path d="M14 2v6h6M8 13h8M8 17h5"/>',
        "repository": '<circle cx="6" cy="5" r="3"/><circle cx="6" cy="19" r="3"/><circle cx="18" cy="5" r="3"/><path d="M6 8v8M18 8a11 11 0 0 1-9 11"/>',
        "pull-request": '<circle cx="6" cy="5" r="3"/><circle cx="6" cy="19" r="3"/><circle cx="18" cy="19" r="3"/><path d="M6 8v8M18 16V9a4 4 0 0 0-4-4h-2M15 2l-3 3 3 3"/>',
        "check": '<path d="m5 12 4 4L19 6"/>',
    }
    return (f'<g transform="translate({x} {y}) scale({size / 24})" fill="none" '
            f'stroke="{PALETTE["pine"]}" stroke-width="1.65" stroke-linecap="round" '
            f'stroke-linejoin="round">{paths[name]}</g>')


def solution_diagram():
    pine, honeydew, shadow = (PALETTE[key] for key in ("pine", "honeydew", "shadow"))
    # A shared brand heading and aligned open columns group both capabilities.
    # Keep the composition free of outer frames, header fills and branching lines.
    artwork = f'''
      <path d="M132 371 H744 M1176 371 H1788" fill="none" stroke="{pine}" stroke-opacity=".18" stroke-width="1.5"/>
      <!-- The report remains an open chart with no surrounding panel. -->
      <g stroke="none">
        <rect x="332" y="670" width="456" height="12" rx="6" fill="{honeydew}"/>
        <rect x="332" y="726" width="456" height="12" rx="6" fill="{honeydew}"/>
        <rect x="332" y="670" width="240" height="12" rx="6" fill="{pine}"/>
        <rect x="332" y="726" width="353" height="12" rx="6" fill="{pine}" opacity=".40"/>
      </g>
      <!-- Monospaced diff, with a shared gutter and restrained fills. -->
      <path d="M1120 674 H1788" stroke="{pine}" stroke-opacity=".12" stroke-width="1.5"/>
      <rect x="1120" y="690" width="668" height="45" rx="5" fill="{shadow}" opacity=".035"/>
      <rect x="1120" y="744" width="668" height="45" rx="5" fill="{honeydew}" opacity=".85"/>
      <g fill="none" stroke="{pine}" stroke-width="1.8" stroke-linecap="round">
        <path d="M1138 712 H1150 M1138 766 H1150 M1144 760 V772" stroke-opacity=".65"/>
        <path d="M1162 699 V725 M1162 753 V779" stroke-opacity=".12" stroke-width="1"/>
      </g>
    '''
    artwork += solution_icon("brief", 132, 535, 28)
    artwork += solution_icon("repository", 1120, 535, 28)
    artwork += solution_icon("responses", 132, 804, 28)
    artwork += solution_icon("sources", 362, 804, 28)
    artwork += solution_icon("evidence", 576, 804, 28)
    artwork += solution_icon("check", 1120, 804, 28)
    artwork += solution_icon("brief", 132, 901, 30)
    artwork += solution_icon("pull-request", 1120, 901, 30)
    return svg_document(artwork, title="Fortop: visibilidad de marca y remediación SEO/GEO en un mismo producto")


def build_solution(deck):
    assets = PITCH / "assets" / "04-solucion"
    assets.mkdir(parents=True, exist_ok=True)
    diagram = solution_diagram()
    (assets / "dos-recorridos.svg").write_text(diagram)
    symbol_group = group_from_asset("01_ISOTIPO/FORTOP_isotipo_color-transparente.svg")
    symbol = f'<svg xmlns="{SVG_NS}" viewBox="24 24 224 192" width="224" height="192">{symbol_group}</svg>'
    center_logo = logo_asset()
    (assets / "fortop-central.svg").write_text(center_logo)
    logo_width = 350
    logo_height = logo_width * 192 / 864.8
    logo_x, logo_y = (W - logo_width) / 2, 371 - logo_height / 2

    def label(value, x, baseline, size, color="shadow", width=1656):
        return {"text": value, "x": x, "baseline": baseline, "size": size, "color": color, "weight": 400, "width": width}

    texts = [
        label("LA SOLUCIÓN", 132, 112, 25, "pine"),
        label("Del diagnóstico a la acción.", 132, 232, 68),
        label("Visibilidad de marca", 132, 510, 40, "pine", width=668),
        label("Remediación SEO/GEO", 1120, 510, 40, "pine", width=668),
        label("Sube tu brief", 174, 561, 26, width=626),
        label("Conecta tu repositorio", 1162, 561, 26, width=626),

        label("Informe de visibilidad", 132, 622, 24, "pine", width=668),
        label("Tu marca", 132, 684, 27, width=175),
        label("Competencia", 132, 740, 27, width=175),
        label("Respuestas", 174, 828, 24, "pine", width=160),
        label("Fuentes", 404, 828, 24, "pine", width=146),
        label("Evidencias", 618, 828, 24, "pine", width=170),

        label("Correcciones propuestas", 1120, 622, 24, "pine", width=668),
        label("index.html", 1120, 663, 21, "pine", width=668),
        label("<title>Inicio</title>", 1174, 722, 27, width=614),
        label("<title>Software para pymes</title>", 1174, 776, 27, "pine", width=614),
        label("Comprobaciones superadas", 1162, 828, 26, "pine", width=626),

        label("Informe y plan de acción", 176, 930, 32, width=624),
        label("Issues y pull requests", 1164, 930, 32, width=624),
        label("04 / 08", 1708, 1022, 20, "pine", width=80),
    ]
    for item in texts:
        if item["text"].startswith("<title>"):
            content = item["text"][7:-8]
            item.update(font_family="Consolas", font_path=CODE_FONT_PATH, size=26,
                        segments=[{"text": "<title>", "color": "pine"},
                                  {"text": content, "color": "shadow"},
                                  {"text": "</title>", "color": "pine"}])
    body = f'<rect width="{W}" height="{H}" fill="{PALETTE["ivory"]}"/>'
    body += place(diagram, 0, 0, W, H)
    body += place(center_logo, logo_x, logo_y, logo_width, logo_height)
    body += place(symbol, 1738, 75, 50, 42.86)
    for item in texts:
        outlines, width = text_paths(item)
        if width > item["width"] or item["x"] + width > W - 100:
            raise ValueError(f'Text exceeds its column: {item["text"]}')
        body += outlines
    svg = svg_document(body, title="Fortop · solución")
    (SLIDES / "04-solucion.svg").write_text(svg)
    cairosvg.svg2png(bytestring=svg.encode(), write_to=str(SLIDES / "04-solucion.png"), output_width=W, output_height=H)
    cairosvg.svg2pdf(bytestring=svg.encode(), write_to=str(SLIDES / "04-solucion.pdf"), output_width=1280, output_height=720)

    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string(PALETTE["ivory"][1:])
    add_vector_picture(slide, diagram, "solucion-diagrama", 0, 0, W, H)
    add_vector_picture(slide, center_logo, "solucion-logo-central", logo_x, logo_y, logo_width, logo_height)
    add_vector_picture(slide, symbol, "solucion-isotipo", 1738, 75, 50, 42.86)
    for item in texts:
        add_native_text(slide, item)
    slide.notes_slide.notes_text_frame.text = SOLUTION_NOTES


def agents_diagram():
    ivory, honeydew = PALETTE["ivory"], PALETTE["honeydew"]
    icons = {
        "research": '<circle cx="12" cy="12" r="9"/><path d="M3 12h18M12 3c3 3 4 6 4 9s-1 6-4 9c-3-3-4-6-4-9s1-6 4-9Z"/>',
        "measure": '<path d="M21 14a3 3 0 0 1-3 3H8l-5 4V6a3 3 0 0 1 3-3h12a3 3 0 0 1 3 3Z"/><path d="M7 8h10M7 12h6"/>',
        "evidence": '<path d="M14 2H6a2 2 0 0 0-2 2v16a2 2 0 0 0 2 2h12a2 2 0 0 0 2-2V8Z"/><path d="M14 2v6h6M8 15l3 3 5-6"/>',
        "prioritize": '<path d="m3 6 2 2 3-4M11 6h10M11 12h10M11 18h10"/><circle cx="5" cy="12" r="1"/><circle cx="5" cy="18" r="1"/>',
        "analyze": '<path d="M14 2H6a2 2 0 0 0-2 2v16a2 2 0 0 0 2 2h5M14 2v6h6V8l-6-6Z"/><circle cx="15" cy="15" r="5"/><path d="m19 19 3 3M8 7h2"/>',
        "edit": '<path d="m15 5 4 4M4 20l5-1L21 7a2.8 2.8 0 0 0-4-4L5 15l-1 5ZM14 20h7"/>',
        "validate": '<path d="M12 3 20 6v6c0 5-8 9-8 9s-8-4-8-9V6Z"/><path d="m8 12 3 3 5-6"/>',
        "publish": '<circle cx="6" cy="5" r="3"/><circle cx="6" cy="19" r="3"/><circle cx="18" cy="19" r="3"/><path d="M6 8v8M18 16V9a4 4 0 0 0-4-4h-2M15 2l-3 3 3 3"/>',
    }
    body = f'''
      <g fill="none" stroke="{honeydew}" stroke-linecap="round" stroke-linejoin="round">
        <path d="M132 627 H1788" stroke-opacity=".16" stroke-width="1.5"/>
        <g stroke-width="2" stroke-opacity=".5">
          <path d="M516 440 H764 M754 432 764 440 754 448 M896 440 H1144 M1134 432 1144 440 1134 448 M1276 440 H1524 M1514 432 1524 440 1514 448"/>
          <path d="M516 760 H764 M754 752 764 760 754 768 M896 760 H1144 M1134 752 1144 760 1134 768 M1276 760 H1524 M1514 752 1524 760 1514 768"/>
        </g>
        <path d="M1210 712 V701 Q1210 685 1194 685 H846 Q830 685 830 701 V712 M824 704 830 712 836 704" stroke-width="1.8" stroke-opacity=".5"/>
      </g>
    '''
    for names, y in ((["research", "measure", "evidence", "prioritize"], 440),
                     (["analyze", "edit", "validate", "publish"], 760)):
        for name, x in zip(names, (450, 830, 1210, 1590)):
            body += (f'<g transform="translate({x - 40} {y - 40}) scale({80 / 24})" '
                     f'fill="none" stroke="{ivory}" stroke-width="1.2" stroke-linecap="round" '
                     f'stroke-linejoin="round">{icons[name]}</g>')
    return svg_document(body, title="Dos recorridos de agentes y herramientas, con verificación antes de publicar")


def build_agents(deck):
    assets = PITCH / "assets" / "05-agentes"
    assets.mkdir(parents=True, exist_ok=True)
    diagram = agents_diagram()
    (assets / "recorridos-y-validacion.svg").write_text(diagram)
    symbol_group = group_from_asset("01_ISOTIPO/FORTOP_isotipo_mono-ivory.svg")
    symbol = f'<svg xmlns="{SVG_NS}" viewBox="24 24 224 192" width="224" height="192">{symbol_group}</svg>'

    def label(value, x, baseline, size, color="ivory", width=1656, center=None):
        item = {"text": value, "x": x, "baseline": baseline, "size": size,
                "color": color, "weight": 400, "width": width}
        if center is not None:
            _, measured = text_paths(item)
            item["x"] = center - measured / 2
            item["width"] = min(width, W - 132 - item["x"])
        return item

    texts = [
        label("CÓMO FUNCIONA", 132, 112, 25),
        label("Agentes con herramientas.", 132, 232, 68),
        label("Google Gemini", 132, 292, 28, "honeydew"),
        label("MARCA", 132, 449, 23, "honeydew", width=190),
        label("REPOSITORIO", 132, 769, 23, "honeydew", width=190),
        label("Si falla, corrige", 0, 666, 22, "honeydew", width=310, center=1020),
        label("05 / 08", 1708, 1022, 20, "honeydew", width=80),
    ]
    rows = [
        (548, [("Investiga", "Búsqueda web"), ("Mide", "Pregunta y repite"),
               ("Contrasta", "Citas e identidad"), ("Prioriza", "Acciones con evidencia")]),
        (868, [("Analiza", "Lee y busca código"), ("Corrige", "Edita archivos"),
               ("Verifica", "Sintaxis y alcance"), ("Publica", "Issues y PRs")]),
    ]
    for baseline, stages in rows:
        for x, (name, detail) in zip((450, 830, 1210, 1590), stages):
            texts.append(label(name, 0, baseline, 40, width=300, center=x))
            texts.append(label(detail, 0, baseline + 42, 24, "honeydew", width=300, center=x))

    body = f'<rect width="{W}" height="{H}" fill="{PALETTE["shadow"]}"/>'
    body += place(diagram, 0, 0, W, H)
    body += place(symbol, 1738, 75, 50, 42.86)
    for item in texts:
        outlines, width = text_paths(item)
        if width > item["width"] or item["x"] < 100 or item["x"] + width > W - 100:
            raise ValueError(f'Text exceeds its column: {item["text"]}')
        body += outlines
    svg = svg_document(body, title="Fortop · sistema agéntico")
    (SLIDES / "05-agentes.svg").write_text(svg)
    cairosvg.svg2png(bytestring=svg.encode(), write_to=str(SLIDES / "05-agentes.png"), output_width=W, output_height=H)
    cairosvg.svg2pdf(bytestring=svg.encode(), write_to=str(SLIDES / "05-agentes.pdf"), output_width=1280, output_height=720)

    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string(PALETTE["shadow"][1:])
    add_vector_picture(slide, diagram, "agentes-diagrama", 0, 0, W, H)
    add_vector_picture(slide, symbol, "agentes-isotipo", 1738, 75, 50, 42.86)
    for item in texts:
        add_native_text(slide, item)
    slide.notes_slide.notes_text_frame.text = AGENTS_NOTES


def screenshot_svg(source, crop, x, y, width):
    """Embed the original bitmap through a crop viewport, without image edits."""
    with Image.open(source) as bitmap:
        image_width, image_height = bitmap.size
    left, top, crop_width, crop_height = crop
    if min(left, top) < 0 or left + crop_width > image_width or top + crop_height > image_height:
        raise ValueError(f"Screenshot crop exceeds the original: {source.name}")
    height = width * crop_height / crop_width
    encoded = base64.b64encode(source.read_bytes()).decode()
    return (f'<svg x="{x}" y="{y}" width="{width}" height="{height}" '
            f'viewBox="{left} {top} {crop_width} {crop_height}" overflow="hidden">'
            f'<image xmlns:xlink="http://www.w3.org/1999/xlink" '
            f'width="{image_width}" height="{image_height}" '
            f'xlink:href="data:image/png;base64,{encoded}"/></svg>')


def add_screenshot(slide, source, crop, x, y, width):
    with Image.open(source) as bitmap:
        image_width, image_height = bitmap.size
    left, top, crop_width, crop_height = crop
    picture = slide.shapes.add_picture(str(source), Inches(x / 144), Inches(y / 144),
                                       width=Inches(width / 144),
                                       height=Inches(width * crop_height / crop_width / 144))
    picture.name = source.stem
    picture.crop_left = left / image_width
    picture.crop_top = top / image_height
    picture.crop_right = (image_width - left - crop_width) / image_width
    picture.crop_bottom = (image_height - top - crop_height) / image_height


def build_product(deck):
    assets = PITCH / "assets" / "06-producto"
    assets.mkdir(parents=True, exist_ok=True)
    logo = logo_asset()
    (assets / "imagotipo-demo.svg").write_text(logo)
    logo_width = 900
    logo_height = logo_width * 192 / 864.8
    logo_x = (W - logo_width) / 2

    def label(value, x, baseline, size, color="shadow", width=1656, center=None, href=None):
        item = {"text": value, "x": x, "baseline": baseline, "size": size,
                "color": color, "weight": 400, "width": width}
        if center is not None:
            _, measured = text_paths(item)
            item["x"] = center - measured / 2
        if href:
            item["href"] = href
        return item

    button_x, button_y, button_width, button_height = 790, 632, 340, 88
    button_text = label("Abrir demo", 0, 688, 32, "ivory", width=220, href=DEMO_URL)
    _, text_width = text_paths(button_text)
    content_width = text_width + 24 + 24
    button_text["x"] = W / 2 - content_width / 2
    arrow_x = button_text["x"] + text_width + 24 - button_x
    button = svg_document(f'''
        <rect width="340" height="88" rx="44" fill="{PALETTE['pine']}"/>
        <path d="M{arrow_x + 3} 54 L{arrow_x + 21} 36 M{arrow_x + 4} 36 H{arrow_x + 21} V53"
              fill="none" stroke="{PALETTE['ivory']}" stroke-width="2.2" stroke-linecap="round" stroke-linejoin="round"/>
    ''', width=button_width, height=button_height, title="Abrir la demo de Fortop")
    (assets / "abrir-demo.svg").write_text(button)
    texts = [
        label("DEMO", 132, 112, 25, "pine"),
        button_text,
        label("06 / 08", 1708, 1022, 20, "pine", width=80),
    ]
    body = f'<rect width="{W}" height="{H}" fill="{PALETTE["ivory"]}"/>'
    body += place(logo, logo_x, 300, logo_width, logo_height)
    body += f'<a href="{DEMO_URL}" target="_blank">{place(button, button_x, button_y, button_width, button_height)}</a>'
    for item in texts:
        outlines, width = text_paths(item)
        if width > item["width"]:
            raise ValueError(f'Text exceeds its column: {item["text"]}')
        body += f'<a href="{item["href"]}" target="_blank">{outlines}</a>' if item.get("href") else outlines
    svg = svg_document(body, title="Fortop · demo en directo")
    (SLIDES / "06-producto.svg").write_text(svg)
    cairosvg.svg2png(bytestring=svg.encode(), write_to=str(SLIDES / "06-producto.png"), output_width=W, output_height=H)
    cairosvg.svg2pdf(bytestring=svg.encode(), write_to=str(SLIDES / "06-producto.pdf"), output_width=1280, output_height=720)

    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string(PALETTE["ivory"][1:])
    add_vector_picture(slide, logo, "demo-logo", logo_x, 300, logo_width, logo_height)
    link_button = add_vector_picture(slide, button, "demo-abrir", button_x, button_y, button_width, button_height)
    link_button.click_action.hyperlink.address = DEMO_URL
    for item in texts:
        add_native_text(slide, item)
        if item.get("href"):
            slide.shapes[-1].click_action.hyperlink.address = item["href"]
    slide.notes_slide.notes_text_frame.text = PRODUCT_NOTES


def business_diagram():
    ivory, honeydew = PALETTE["ivory"], PALETTE["honeydew"]
    body = f'''
      <g fill="none" stroke="{honeydew}" stroke-width="1.8" stroke-opacity=".48" stroke-linecap="round" stroke-linejoin="round">
        <path d="M622 497 H722 M710 489 L722 497 710 505 M1200 497 H1300 M1288 489 L1300 497 1288 505"/>
        <path d="M132 840 H1788" stroke-opacity=".25" stroke-width="1.3"/>
      </g>
      <!-- Buyer: the founder of a small software business. -->
      <g fill="none" stroke="{ivory}" stroke-width="3" stroke-linecap="round" stroke-linejoin="round">
        <circle cx="382" cy="467" r="20"/>
        <path d="M344 540 V526 C344 494 420 494 420 526 V540"/>
        <!-- Service: a report and a code change, on one document. -->
        <path d="M927 448 H971 L995 472 V543 H927 Z"/>
        <path d="M971 448 V472 H995 M944 484 H979" stroke-opacity=".65"/>
        <path d="M947 500 L937 512 947 524 M976 500 L986 512 976 524 M966 497 L956 528"/>
        <!-- Payment: monthly billing. The euro sign is editable text. -->
        <rect x="1494" y="453" width="88" height="88" rx="8"/>
        <path d="M1513 443 V465 M1563 443 V465 M1494 477 H1582"/>
      </g>
    '''
    return svg_document(body, title="Quién paga, qué recibe y cómo se cobra")


def build_business(deck):
    assets = PITCH / "assets" / "07-negocio"
    assets.mkdir(parents=True, exist_ok=True)
    diagram = business_diagram()
    (assets / "cliente-y-suscripcion.svg").write_text(diagram)
    symbol_group = group_from_asset("01_ISOTIPO/FORTOP_isotipo_mono-ivory.svg")
    symbol = f'<svg xmlns="{SVG_NS}" viewBox="24 24 224 192" width="224" height="192">{symbol_group}</svg>'

    def label(value, x, baseline, size, color="ivory", width=1656, center=None):
        item = {"text": value, "x": x, "baseline": baseline, "size": size,
                "color": color, "weight": 400, "width": width}
        if center is not None:
            _, measured = text_paths(item)
            item["x"] = center - measured / 2
            item["width"] = min(width, W - 132 - item["x"])
        return item

    texts = [
        label("CLIENTE Y NEGOCIO", 132, 112, 25),
        label("Para empresas de software", 132, 232, 68),
        label("sin equipo SEO.", 132, 306, 68),
        label("QUIÉN PAGA", 0, 383, 24, "honeydew", width=500, center=382),
        label("QUÉ RECIBE", 0, 383, 24, "honeydew", width=500, center=960),
        label("CÓMO COBRAMOS", 0, 383, 24, "honeydew", width=500, center=1538),
        label("Fundador", 0, 625, 44, width=500, center=382),
        label("de un SaaS pequeño", 0, 675, 28, "honeydew", width=500, center=382),
        label("Con web en GitHub", 0, 728, 27, "honeydew", width=500, center=382),
        label("Diagnóstico de marca", 0, 625, 38, width=500, center=960),
        label("y correcciones web", 0, 675, 40, width=500, center=960),
        label("Cambios para revisar", 0, 728, 27, "honeydew", width=500, center=960),
        label("Cuota mensual", 0, 625, 44, width=500, center=1538),
        label("Según volumen de análisis", 0, 675, 27, "honeydew", width=500, center=1538),
        label("€", 0, 524, 45, width=60, center=1538),
        label("Primero: probarlo con fundadores de pequeños SaaS.", 132, 933, 34),
        label("07 / 08", 1708, 1022, 20, "honeydew", width=80),
    ]
    body = f'<rect width="{W}" height="{H}" fill="{PALETTE["pine"]}"/>'
    body += place(diagram, 0, 0, W, H)
    body += place(symbol, 1738, 75, 50, 42.86)
    for item in texts:
        outlines, width = text_paths(item)
        if width > item["width"] or item["x"] < 100 or item["x"] + width > W - 100:
            raise ValueError(f'Text exceeds its column: {item["text"]}')
        body += outlines
    svg = svg_document(body, title="Fortop · cliente y modelo de negocio")
    (SLIDES / "07-negocio.svg").write_text(svg)
    cairosvg.svg2png(bytestring=svg.encode(), write_to=str(SLIDES / "07-negocio.png"), output_width=W, output_height=H)
    cairosvg.svg2pdf(bytestring=svg.encode(), write_to=str(SLIDES / "07-negocio.pdf"), output_width=1280, output_height=720)

    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string(PALETTE["pine"][1:])
    add_vector_picture(slide, diagram, "negocio-diagrama", 0, 0, W, H)
    add_vector_picture(slide, symbol, "negocio-isotipo", 1738, 75, 50, 42.86)
    for item in texts:
        add_native_text(slide, item)
    slide.notes_slide.notes_text_frame.text = BUSINESS_NOTES


def build_closing(deck):
    assets = PITCH / "assets" / "08-cierre"
    assets.mkdir(parents=True, exist_ok=True)
    logo = logo_asset()
    (assets / "imagotipo-original.svg").write_text(logo)

    def label(value, baseline, size, color="shadow", width=1656, center=960):
        item = {"text": value, "x": 0, "baseline": baseline, "size": size,
                "color": color, "weight": 400, "width": width}
        _, measured = text_paths(item)
        item["x"] = center - measured / 2
        item["width"] = min(width, W - 132 - item["x"])
        return item

    texts = [
        label("Entiende tu visibilidad.", 570, 64),
        label("Mejora tu web.", 650, 64, "pine"),
        {"text": TEAM_CREDITS, "x": 132, "baseline": 1022, "size": 22,
         "color": "pine", "weight": 400, "width": 1480},
        {"text": "08 / 08", "x": 1708, "baseline": 1022, "size": 20,
         "color": "pine", "weight": 400, "width": 80},
    ]
    logo_width = 850
    logo_height = logo_width * 192 / 864.8
    logo_x = (W - logo_width) / 2
    body = f'<rect width="{W}" height="{H}" fill="{PALETTE["ivory"]}"/>'
    body += place(logo, logo_x, 220, logo_width, logo_height)
    for item in texts:
        outlines, width = text_paths(item)
        if width > item["width"] or item["x"] < 100 or item["x"] + width > W - 100:
            raise ValueError(f'Text exceeds its column: {item["text"]}')
        body += outlines
    svg = svg_document(body, title="Fortop · cierre")
    (SLIDES / "08-cierre.svg").write_text(svg)
    cairosvg.svg2png(bytestring=svg.encode(), write_to=str(SLIDES / "08-cierre.png"), output_width=W, output_height=H)
    cairosvg.svg2pdf(bytestring=svg.encode(), write_to=str(SLIDES / "08-cierre.pdf"), output_width=1280, output_height=720)

    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string(PALETTE["ivory"][1:])
    add_vector_picture(slide, logo, "cierre-imagotipo", logo_x, 220, logo_width, logo_height)
    for item in texts:
        add_native_text(slide, item)
    slide.notes_slide.notes_text_frame.text = CLOSING_NOTES


def main():
    for folder in [SLIDES, ASSETS]:
        folder.mkdir(parents=True, exist_ok=True)
    if not FONT_PATH.exists():
        raise FileNotFoundError("Set FORTOP_PITCH_FONT to the Bahnschrift font installation.")

    logo, art = logo_asset(), ascenso_art()
    (ASSETS / "imagotipo-original.svg").write_text(logo)
    (ASSETS / "ascenso-relieve.svg").write_text(art)
    body = f'<rect width="{W}" height="{H}" fill="{PALETTE["ivory"]}"/>'
    body += place(art, 1100, 174, 860, 760)
    body += place(logo, 132, 338, 900, 199.82)
    for item in TEXT:
        outlines, width = text_paths(item)
        if width > item.get("width", 1080 - item["x"]):
            raise ValueError(f'Text exceeds its cover area: {item["text"]}')
        body += outlines
    svg = svg_document(body)
    (SLIDES / "01-portada.svg").write_text(svg)
    cairosvg.svg2png(bytestring=svg.encode(), write_to=str(SLIDES / "01-portada.png"), output_width=W, output_height=H)
    cairosvg.svg2pdf(bytestring=svg.encode(), write_to=str(SLIDES / "01-portada.pdf"), output_width=1280, output_height=720)

    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(W / 144), Inches(H / 144)
    deck.core_properties.title = "Fortop · Pitch"
    deck.core_properties.subject = "Visibilidad de marca y remediación SEO/GEO"
    deck.core_properties.author = "; ".join(TEAM_MEMBERS)
    deck.core_properties.language = "es-ES"
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string(PALETTE["ivory"][1:])
    add_vector_picture(slide, art, "ascenso-relieve", 1100, 174, 860, 760)
    add_vector_picture(slide, logo, "imagotipo-original", 132, 338, 900, 199.82)
    for item in TEXT:
        add_native_text(slide, item)
    slide.notes_slide.notes_text_frame.text = NOTES
    build_problem(deck)
    build_insight(deck)
    build_solution(deck)
    build_agents(deck)
    build_product(deck)
    build_business(deck)
    build_closing(deck)
    deck.save(PITCH / "Fortop-Pitch.pptx")
    print("Built all 8 slides: SVG, PNG, PDF and editable PowerPoint.")


if __name__ == "__main__":
    main()
